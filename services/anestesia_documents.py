"""Company-authored DOCX documents. Issued certificates are never reconstructed."""
from __future__ import annotations

from datetime import date
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from services.anestesia_docs import totals
from services import lp_documents as _lp
if getattr(_lp, "LP_DOCUMENTS_API_VERSION", 0) < 2:
    import importlib
    _lp = importlib.reload(_lp)
_add_company_header = _lp._add_company_header
get_lp_company_profile = _lp.get_lp_company_profile
render_lp_document = _lp.render_lp_document

MONTHS = ("", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre")


def rir_profile():
    # Verified against Registro Público 15-May-2026 and the signed Bella Vista
    # operations notice. Keep the other LP company profiles untouched.
    return replace(get_lp_company_profile("RIR Medical"), legal_name="RIR MEDICAL ENGINEERING, S.EP.",
        operations_notice="155750585-2-2024-2024-574365876",
        address="PH Bonanza Plaza, local 4B, calle 41, Bella Vista, Panamá",
        header_lines=("RUC: 155750585-2-2024 | DV: 40", "PH Bonanza Plaza, local 4B · Bella Vista, Panamá", "Tel. +507 6847-5616"))


def quote_docx(source: dict, config: dict, assets: Path) -> bytes:
    """US Letter business brief, existing RIR branding, fixed-width price table."""
    profile = rir_profile()
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin, section.bottom_margin = Inches(1.35), Inches(0.65)  # branded-header override
        section.left_margin = section.right_margin = Inches(1)
        section.header_distance = section.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_before, normal.paragraph_format.space_after = Pt(0), Pt(4)
    normal.paragraph_format.line_spacing = 1.04
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Title"].paragraph_format.space_after = Pt(8)
    for key, size in (("Heading 1", 16), ("Heading 2", 12)):
        style = doc.styles[key]
        style.font.name, style.font.size, style.font.color.rgb = "Calibri", Pt(size), RGBColor.from_string("173C55")
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(10), Pt(6)
    _add_company_header(doc, profile, assets)
    issued = date.fromisoformat(config["document_date"])
    doc.add_heading("COTIZACIÓN", 0)
    doc.add_paragraph(f"{issued.day} de {MONTHS[issued.month]} de {issued.year}")
    for label, value in (("Acto", source["number"]), ("Dirigida a", source["entity"]),
                         ("Unidad de compra", source["purchase_unit"]), ("Objeto", source["title"])):
        p = doc.add_paragraph(); p.add_run(label + ": ").bold = True; p.add_run(value)
    doc.add_paragraph(f"{profile.legal_name}, RUC {profile.ruc}, DV {profile.dv}, presenta la siguiente oferta:")
    item = source["items"][0]
    amounts = totals(item["cantidad"], config["price"], config["tax_mode"], config.get("tax_rate", 7))
    table = doc.add_table(rows=1, cols=4)
    table.style, table.autofit = "Table Grid", False
    widths = [1200, 4140, 2010, 2010]
    tbl_pr = table._tbl.tblPr
    for name, attrs in (("tblW", {"w": "9360", "type": "dxa"}), ("tblInd", {"w": "120", "type": "dxa"})):
        old = tbl_pr.find(qn("w:" + name))
        if old is not None: tbl_pr.remove(old)
        el = OxmlElement("w:" + name)
        for k, v in attrs.items(): el.set(qn("w:" + k), v)
        tbl_pr.append(el)
    for cell, text in zip(table.rows[0].cells, ("Cantidad", "Producto", "Precio unitario USD", "Importe USD")):
        cell.text = text
        for run in cell.paragraphs[0].runs: run.bold = True
        shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "EAF1F5"); cell._tc.get_or_add_tcPr().append(shade)
    cells = table.add_row().cells
    cells[0].text = amounts["cantidad"]
    cells[1].text = (f"Kit de circuito de paciente para máquina de anestesia.\nFicha CTNI 43358.\n"
                     f"Marca: {config['catalog_brand']}. Modelo/catálogo: {config['catalog_model']}.\n"
                     f"Presentación: {item.get('unidad') or 'Unidad'}.")
    unit_price = Decimal(amounts["precio_ingresado"])
    digits = max(2, -unit_price.as_tuple().exponent)
    cells[2].text = f"{unit_price:,.{digits}f}"
    cells[3].text = f"{float(amounts['total'] if config['tax_mode'] == 'incluido' else amounts['subtotal']):,.2f}"
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width / 1440)
        for row in table.rows:
            cell = row.cells[i]
            cell.width = Inches(width / 1440)
            cell._tc.get_or_add_tcPr().get_or_add_tcW().set(qn("w:w"), str(width))
    for key, label in (("subtotal", "Subtotal"), ("itbms", "ITBMS"), ("total", "TOTAL OFERTADO")):
        p = doc.add_paragraph(f"{label}: USD {float(amounts[key]):,.2f}")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if key == "total":
            for run in p.runs: run.bold = True
    mode = {"exento": "Exento / no aplica ITBMS", "incluido": "El precio unitario ingresado incluye ITBMS", "adicional": "ITBMS adicional al precio unitario"}[config["tax_mode"]]
    doc.add_paragraph(mode + (f" ({amounts['tasa']}%)." if config["tax_mode"] != "exento" else "."))
    doc.add_heading("Condiciones de la oferta", level=2)
    for label, value in (("Lugar de entrega", config["delivery_place"]), ("Entregas", config["delivery"]),
                         ("Forma de pago", source["info"].get("forma de pago", "Crédito")),
                         ("Validez de la cotización", f"{config.get('proposal_validity_days', 30)} días calendario"),
                         ("Garantía / vida útil exigida", config.get("warranty", ""))):
        if value:
            p = doc.add_paragraph(); p.add_run(label + ": ").bold = True; p.add_run(str(value))
    p = doc.add_paragraph("Atentamente,"); p.paragraph_format.keep_with_next = True
    signature = assets / profile.signature_filename
    if not signature.is_file(): raise ValueError("Falta la firma autorizada de RIR.")
    p = doc.add_paragraph(); p.add_run().add_picture(str(signature), width=Inches(1.0)); p.paragraph_format.keep_with_next = True
    doc.add_paragraph(f"{profile.representative}\nRepresentante legal\n{profile.legal_name}")
    output = BytesIO(); doc.save(output)
    return output.getvalue()


def pact_docx(source, config, root):
    day = date.fromisoformat(config["document_date"])
    profile = rir_profile()
    replacements = {"[Representante_legal_de_la_Entidad_Licitante]": config["entity_representative"],
        "[cargo_entidad]": config["entity_role"], "[cedula]": config["entity_id"], "[entidad]": source["entity"],
        "[numero_de_acto]": source["number"], "[titulo]": source["title"], "[lugar]": config["delivery_place"],
        "[entrega]": config["delivery"], "[dia]": str(day.day), "[mes]": MONTHS[day.month], "[año]": str(day.year),
        "[fecha]": f"{day.day} de {MONTHS[day.month]} de {day.year}",
        "[identidad_contratista]": (f"{profile.representative}, con cédula de identidad personal No. {profile.representative_id}, "
            f"actuando en nombre y representación de {profile.legal_name}, sociedad de emprendimiento inscrita "
            f"en el Registro Público a folio 155750585, con RUC {profile.ruc}, DV {profile.dv}, "
            f"Aviso de Operación No. {profile.operations_notice}, con domicilio en {profile.address}")}
    return render_lp_document(root / "assets/doc_gen_base/template_pacto_de_integridad_rir.docx", replacements,
        company_name="RIR Medical", document_name="pacto_de_integridad.docx", assets_dir=root / "assets/cotizacion_base", company_profile=profile)
