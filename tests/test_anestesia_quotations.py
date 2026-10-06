"""Quotation-only output: official data, totals, retries and act isolation."""
from copy import deepcopy
from datetime import timedelta
from io import BytesIO
import json
from unittest.mock import patch

from docx import Document
import fitz
import pytest
from streamlit.testing.v1 import AppTest

from services import anestesia_quotations as quotes, anestesia_view as view
from services.anestesia_docs import PANAMA, file_hash
from services.anestesia_storage import AnestesiaStorage
from services.anestesia_worker import run_request
from test_anestesia_docs import fixture, FakeStorage, ROOT, IDENT, ACT, URL
from test_anestesia_delivery import Storage as DriveStorage, contents
from test_anestesia_view import Storage as ViewStorage

ACT2 = "2026-1-10-01-06-CL-051243"
URL2 = "https://www.panamacompra.gob.pa/Inicio/#/solicitud-de-cotizacion/" + ACT2 + "/0nM6ICc0JCLwUDN3QDMxojIpJye"


class Storage(DriveStorage):
    ensure_quotation = AnestesiaStorage.ensure_quotation
    save_quotation = AnestesiaStorage.save_quotation
    quotation_root = AnestesiaStorage.quotation_root
    trash_file = AnestesiaStorage.trash_file

    def __init__(self):
        super().__init__()
        source, cfg, library = fixture()
        source["info"].update(provincia="Panamá", **{"termino de entrega": "30 Días hábiles"})
        self.source = source
        originals = FakeStorage(source, library)
        docs = []
        for doc in originals.tables["ANESTESIA_DOCUMENTOS"]:
            file = self.put("library", doc["kind"] + ".pdf", originals.get_bytes(doc["file_id"]), "application/pdf")
            docs.append({**doc, **file})
        self.tables = {"ANESTESIA_COTIZACIONES": [], "ANESTESIA_DOCUMENTOS": docs}
        self.jobs = {IDENT: {"id": IDENT, "number": ACT, "url": URL, "state": "Nuevo"}}
        self.fail_conversion = False

    def rows(self, table): return deepcopy(self.tables[table])
    def job(self, ident): return deepcopy(self.jobs.get(ident))
    def save_job(self, data):
        self.jobs[data["id"]].update(deepcopy(data))
        return self.job(data["id"])
    def _write(self, name, columns, data, row=None):
        saved = deepcopy(data)
        if row:
            saved["_row"] = row
            self.tables[name][row - 2] = saved
        else:
            saved["_row"] = len(self.tables[name]) + 2
            self.tables[name].append(saved)
    def put(self, *args):
        file = super().put(*args)
        file["url"] = "https://drive.google.com/file/d/" + file["file_id"] + "/view"
        return file
    def convert_document(self, data, name, parent):
        if self.fail_conversion:
            raise TimeoutError("Conversión interrumpida")
        doc = Document(BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        with fitz.open() as pdf:
            page = pdf.new_page(width=612, height=1008)
            page.insert_textbox((30, 30, 580, 780), text, fontsize=9)
            return pdf.tobytes()


def values():
    return {"catalog": "K", "price": "19.48", "tax_mode": "exento", "tax_rate": 7,
        "warehouse": "Almacén médico quirúrgico", "location_reviewed": True,
        "delivery_use_portal": True, "workflow": "quotation_v2"}


def generate(storage, monkeypatch, execution="exec-1", ident=IDENT, source=None, cfg=None):
    source = deepcopy(source or storage.source)
    monkeypatch.setattr(quotes, "capture", lambda *a, **kw: deepcopy(source))
    return run_request(storage, {"request_id": ident, "action": "generate_quotation",
        "url": source["url"], "config": cfg or values()}, execution_id=execution, root=ROOT)


def test_one_button_generates_only_quotation_pdf_and_word_and_registers_amounts_and_official_date(monkeypatch):
    storage = Storage()
    before = set(storage.api.files_data)
    result = generate(storage, monkeypatch)
    assert result["state"] == "Documentos generados"
    assert result["quotation_number"] == ACT
    record = storage.rows("ANESTESIA_COTIZACIONES")[0]
    assert record["config"]["document_date"] == storage.source["publication"]
    assert record["config"]["proposal_validity_days"] == 120
    assert record["config"]["delivery_place"] == "Panamá - Ciudad de la Salud - Almacén médico quirúrgico"
    assert record["amounts"]["total"] == "17532.00"
    live = contents(storage, result["folder_id"])
    basename = quotes.quotation_filename(storage.source)
    assert set(live) == {basename + ".pdf", basename + ".docx"}
    assert result["delivery_pdf_count"] == record["delivery_pdf_count"] == 1
    assert result["quotation_output_version"] == record["quotation_output_version"] == 3
    assert record["word_url"] == result["word_url"] and record["final_url"] == result["final_url"]
    generated = [f for ident, f in storage.api.files_data.items() if ident not in before]
    assert not any(f["name"].endswith(".zip") for f in generated)
    assert all(f["name"] == basename + ".pdf" for f in generated if f["mimeType"] == "application/pdf")
    assert sum(f["name"] == basename + ".docx" for f in generated) == 2  # staged source and verified final copy
    assert result["zip_url"] == record["zip_url"] == ""
    assert storage.api.files_data[result["folder_id"]]["name"] == "Cotizaciones generadas"
    assert result["source_preview"]["fingerprint"] == storage.source["fingerprint"]
    assert result["source_id"] == ""
    assert len(live) == 2  # no manifests, annexes or working folders in the final folder


def test_new_act_gets_next_number_while_same_act_regenerates_without_overwriting_other_act(monkeypatch):
    storage = Storage()
    first = generate(storage, monkeypatch)
    old = contents(storage, first["folder_id"])
    second_source = {**storage.source, "number": ACT2, "url": URL2, "fingerprint": "official-2"}
    ident2 = "b" * 32
    storage.jobs[ident2] = {"id": ident2, "number": ACT2, "state": "Nuevo"}
    second = generate(storage, monkeypatch, "exec-2", ident2, second_source)
    assert second["quotation_number"] == ACT2
    assert first["folder_id"] == second["folder_id"]
    both = contents(storage, first["folder_id"])
    assert len(both) == 4 and all(both[name] == data for name, data in old.items())
    updated = generate(storage, monkeypatch, "exec-3", cfg={**values(), "price": "21"})
    assert updated["quotation_number"] == ACT
    assert updated["folder_id"] == first["folder_id"]
    assert len(storage.rows("ANESTESIA_COTIZACIONES")) == 2
    replaced = contents(storage, first["folder_id"])
    assert len(replaced) == 4
    assert any(replaced[name] != data for name, data in old.items())
    assert all(replaced[name] == data for name, data in both.items() if ACT2 in name)


def test_retry_after_conversion_failure_preserves_consecutive_and_successful_execution_is_idempotent(monkeypatch):
    storage = Storage()
    storage.fail_conversion = True
    with pytest.raises(TimeoutError): generate(storage, monkeypatch)
    assert storage.rows("ANESTESIA_COTIZACIONES")[0]["quotation_number"] == ACT
    storage.fail_conversion = False
    result = generate(storage, monkeypatch)
    before = len(storage.api.files_data)
    assert generate(storage, monkeypatch) == result
    assert len(storage.api.files_data) == before
    assert len(storage.rows("ANESTESIA_COTIZACIONES")) == 1


@pytest.mark.parametrize("kind", ["css", "dgi", "registro_publico", "oferente", "criterio_tecnico"])
def test_expired_certificate_does_not_block_a_quotation(monkeypatch, kind):
    storage = Storage()
    for row in storage.tables["ANESTESIA_DOCUMENTOS"]:
        if row["kind"] == kind:
            row["expires"] = "2000-01-01"
            if kind == "registro_publico": row["issued"] = "2000-01-01"
    result = generate(storage, monkeypatch)
    assert result["state"] == "Documentos generados"
    assert len(contents(storage, result["folder_id"])) == 2


@pytest.mark.parametrize("field,value", [("publication", ""), ("publication", "2099-01-01"),
    ("purchase_unit", ""), ("info", {}), ("explicit_fichas", ["12345"]), ("items", [])])
def test_incomplete_or_wrong_official_data_never_creates_quote(monkeypatch, field, value):
    storage = Storage()
    source = {**storage.source, field: value}
    try:
        result = generate(storage, monkeypatch, source=source)
        assert result["state"] == "Bloqueado"
    except ValueError:
        assert storage.job(IDENT)["state"] == "Bloqueado"
    assert not storage.rows("ANESTESIA_COTIZACIONES")


def test_original_changes_are_irrelevant_to_quotation_and_both_mask_models_work(monkeypatch):
    storage = Storage()
    ct = next(d for d in storage.tables["ANESTESIA_DOCUMENTOS"] if d["kind"] == "criterio_tecnico")
    storage.api.content[ct["file_id"]] += b"changed"
    result = generate(storage, monkeypatch)
    assert result["state"] == "Documentos generados"
    storage = Storage()
    result = generate(storage, monkeypatch, cfg={**values(), "catalog": "C"})
    assert result["state"] == "Documentos generados"
    assert result["config"]["catalog_model"] == "LB4330C"


@pytest.mark.parametrize("cfg", [{"price": "0"}, {"price": "NaN"}, {"price": "-5"},
    {"tax_mode": "unknown"}, {"catalog": "X"}, {"delivery_use_portal": False, "delivery_manual": ""}])
def test_invalid_quotation_inputs_still_block_before_allocating_a_number(monkeypatch, cfg):
    storage = Storage()
    result = generate(storage, monkeypatch, cfg={**values(), **cfg})
    assert result["state"] == "Bloqueado"
    assert any(c["documento"] == "Datos de la cotización" for c in result["checks"])
    assert not storage.rows("ANESTESIA_COTIZACIONES")


@pytest.mark.parametrize("mode", ["empty", "unavailable"])
def test_no_library_is_required_for_a_quotation(monkeypatch, mode):
    storage = Storage()
    storage.tables["ANESTESIA_DOCUMENTOS"] = []
    rows = storage.rows
    def quotation_rows(table):
        if table == "ANESTESIA_DOCUMENTOS":
            pytest.fail("Quotation-only generation must never read the certificate library")
        return rows(table)
    if mode == "unavailable":
        monkeypatch.setattr(storage, "rows", quotation_rows)
    result = generate(storage, monkeypatch)
    assert result["state"] == "Documentos generados"


def test_portal_term_and_delivery_attestations_are_bound_to_fresh_source():
    storage = Storage()
    with pytest.raises(ValueError, match="cambió"):
        quotes.prepare_quotation(storage.source, {**values(), "delivery_portal_source_hash": "old"})
    with pytest.raises(ValueError, match="Verifica"):
        quotes.prepare_quotation(storage.source, {**values(), "warehouse": "", "location_reviewed": False})
    cfg = quotes.prepare_quotation(storage.source, {**values(), "delivery_use_portal": False,
        "delivery_manual": "10 kits al mes durante 3 meses"})
    assert cfg["delivery"] == "10 kits al mes durante 3 meses"


def test_url_alias_keeps_quote_number_and_duplicate_number_corruption_is_rejected():
    storage = Storage()
    a = storage.ensure_quotation(ACT, URL)
    b = storage.ensure_quotation(ACT, URL.replace("www.panamacompra", "panamacompra"))
    assert a["id"] == b["id"] and len(storage.rows("ANESTESIA_COTIZACIONES")) == 1
    storage.tables["ANESTESIA_COTIZACIONES"].append({**a, "id": "duplicate"})
    with pytest.raises(ValueError, match="duplicados"): storage.ensure_quotation(ACT2, URL2)


def test_quote_template_preserves_description_precise_prices_and_all_fixed_conditions():
    from services.anestesia_documents import quote_docx
    storage = Storage()
    source = deepcopy(storage.source)
    source["items"][0]["descripcion"] = "<p>Kit circuito</p><p>Tubo de 182 cm y filtro HMEF</p>"
    cfg = quotes.prepare_quotation(source, values())
    cfg["quotation_number"] = "RIR-000012"
    doc = Document(BytesIO(quote_docx(source, cfg, ROOT / "assets/cotizacion_base")))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for part in ("Descripción del producto", "Kit circuito\nTubo de 182 cm y filtro HMEF", "Ficha técnica: 43358",
        "LB4330K", "MFLAB", "NINGBO MFLAB", "País de origen: China", "País de procedencia: China",
        "Trae impreso y visible la fecha de manufactura", "aseguramiento de calidad y comercialización",
        "120 días calendario", "Forma de pago: Crédito", "Garantía / Vencimiento de la esterilidad",
        "24 meses de garantía y esterilidad no menor a 24 meses", "19.48", "17,532.00"):
        assert part in text
    assert "RIR-000012" not in text and "Número de cotización:" not in text
    assert doc.sections[0].page_height.inches == 14
    assert doc.sections[0].page_width.inches == 8.5
    header = doc.sections[0].header._element.xml
    assert "ENGINEERING" in header and "info@rirmedical.com" in header


def test_portal_standard_observations_are_replaced_once_without_dropping_other_notes():
    from services.anestesia_documents import quote_docx
    storage = Storage()
    source = deepcopy(storage.source)
    source["items"][0]["descripcion"] = "Kit completo.\nOBSERVACIÓN:\n1.Debe traer impreso y visible fecha de manufactura.\n2.Cumplir con los estándares internacionales."
    cfg = quotes.prepare_quotation(source, values())
    doc = Document(BytesIO(quote_docx(source, cfg, ROOT / "assets/cotizacion_base")))
    text = doc.tables[0].cell(1, 1).text
    assert "Debe traer" not in text and "Cumplir con" not in text
    assert text.count("Trae impreso y visible") == 1
    assert text.count("Cumple con los estándares internacionales") == 1
    source["items"][0]["descripcion"] += "\n3. Incluir accesorio específico solicitado."
    doc = Document(BytesIO(quote_docx(source, cfg, ROOT / "assets/cotizacion_base")))
    text = doc.tables[0].cell(1, 1).text
    assert "accesorio específico solicitado" in text and "Debe traer" not in text


def test_source_change_during_generation_never_publishes_new_files(monkeypatch):
    storage = Storage()
    captures = iter([deepcopy(storage.source), {**storage.source, "fingerprint": "new-amendment"}])
    monkeypatch.setattr(quotes, "capture", lambda *a, **kw: next(captures))
    with pytest.raises(ValueError, match="cambiaron durante"):
        run_request(storage, {"request_id": IDENT, "action": "generate_quotation", "url": URL,
            "config": values()}, execution_id="changed-source", root=ROOT)
    assert storage.job(IDENT)["source_hash"] == "new-amendment"
    assert storage.job(IDENT)["state"] == "Bloqueado"
    assert not storage.job(IDENT).get("final_url")


def test_quote_publication_failure_restores_previous_pdfs_and_preserves_number(monkeypatch):
    storage = Storage()
    previous = generate(storage, monkeypatch)
    old = contents(storage, previous["folder_id"])
    storage.api.fail_copy = storage.api.copy_count + 1
    with pytest.raises(RuntimeError, match="restauró"):
        generate(storage, monkeypatch, "failed-copy", cfg={**values(), "price": "20"})
    assert contents(storage, previous["folder_id"]) == old
    assert len(storage.rows("ANESTESIA_COTIZACIONES")) == 1
    storage.api.fail_copy = None
    done = generate(storage, monkeypatch, "retry-copy", cfg={**values(), "price": "20"})
    assert done["quotation_number"] == previous["quotation_number"]
    assert len(contents(storage, done["folder_id"])) == 2


@pytest.mark.parametrize("failure", ["fail_copy", "corrupt_copy"])
def test_second_format_copy_failure_preserves_both_previous_files(monkeypatch, failure):
    storage = Storage()
    previous = generate(storage, monkeypatch)
    old = contents(storage, previous["folder_id"])
    setattr(storage.api, failure, storage.api.copy_count + 2)
    with pytest.raises(RuntimeError, match="restauró"):
        generate(storage, monkeypatch, "failed-second-format", cfg={**values(), "price": "20"})
    assert contents(storage, previous["folder_id"]) == old
    assert storage.rows("ANESTESIA_COTIZACIONES")[0]["publication"] is None


def test_hard_interruption_retains_journal_then_retry_recovers_and_leaves_only_two_files(monkeypatch):
    storage = Storage()
    first = generate(storage, monkeypatch)
    storage.api.crash_copy = storage.api.copy_count + 2
    with pytest.raises(KeyboardInterrupt): generate(storage, monkeypatch, "hard-crash")
    record = storage.rows("ANESTESIA_COTIZACIONES")[0]
    assert record["publication"]
    assert not storage.api.files_data[record["publication"]["stage"]].get("trashed")
    done = generate(storage, monkeypatch, "recover-hard-crash")
    assert done["quotation_number"] == ACT
    assert len(contents(storage, first["folder_id"])) == 2
    assert not storage.rows("ANESTESIA_COTIZACIONES")[0]["publication"]


@pytest.mark.parametrize("pages,size", [(2, (612,1008)), (1, (612,792))])
def test_multi_page_or_letter_pdf_is_not_published(monkeypatch, pages, size):
    storage = Storage()
    def wrong_layout(*args, **kwargs):
        with fitz.open() as pdf:
            for _ in range(pages): pdf.new_page(width=size[0], height=size[1]).insert_text((50,50), ACT)
            return pdf.tobytes()
    monkeypatch.setattr(storage, "convert_document", wrong_layout)
    with pytest.raises(ValueError, match="una sola hoja larga"):
        generate(storage, monkeypatch)
    assert not storage.job(IDENT).get("final_url")


def test_filename_uses_official_recipient_and_act():
    assert quotes.quotation_filename({"entity":"Caja de Seguro Social", "number":ACT}) == "Cotización firmada dirigida a la Caja de Seguro Social - " + ACT
    assert quotes.quotation_filename({"entity":"MINISTERIO DE SALUD", "number":ACT2}) == "Cotización firmada dirigida al Ministerio de Salud - " + ACT2


def test_library_replacement_during_conversion_does_not_affect_quotation(monkeypatch):
    storage = Storage()
    convert = storage.convert_document
    def change(*args, **kw):
        result = convert(*args, **kw)
        storage.tables["ANESTESIA_DOCUMENTOS"][0]["id"] = "new-certificate-version"
        return result
    monkeypatch.setattr(storage, "convert_document", change)
    result = generate(storage, monkeypatch)
    assert result["state"] == "Documentos generados" and result["final_url"]


def test_previous_numbering_changes_to_official_act_without_duplicating_record(monkeypatch):
    storage = Storage()
    quote = storage.ensure_quotation(ACT, URL)
    storage.save_quotation({"id": quote["id"], "quotation_number": "RIR-000001", "folder_id": "old-case-folder"})
    storage.jobs[IDENT].update(state="Documentos generados", last_execution="exec-1",
        delivery_pdf_count=12, zip_url="old-zip")
    result = generate(storage, monkeypatch)
    assert result["quotation_number"] == ACT
    assert result["folder_id"] != "old-case-folder"
    assert len(storage.rows("ANESTESIA_COTIZACIONES")) == 1
    assert len(contents(storage, result["folder_id"])) == 2
    assert result["zip_url"] == ""


APP = "from services.anestesia_view import render_anestesia_docs\nrender_anestesia_docs(None, 'usuario')"


class UIStorage(ViewStorage):
    def __init__(self):
        super().__init__()
        self.tables["ANESTESIA_COTIZACIONES"] = []
        self.tables["ANESTESIA_EXPEDIENTES"] = []
        self.enqueued = []
    def enqueue(self, payload, **kwargs):
        self.enqueued.append(deepcopy(payload))
        return "queue-1"
    def quotation_root(self): return "all-quotes"
    def document_control_links(self): return {"sheet": "https://docs.google.com/spreadsheets/d/control/edit", "folder": "https://drive.google.com/drive/folders/originals"}


@pytest.fixture
def ui():
    view._records.clear(); view._json.clear(); view._live_job.clear(); view._quotation_folder.clear(); view._document_control_links.clear()
    storage = UIStorage()
    with patch.object(view, "AnestesiaStorage", return_value=storage), patch.object(view, "build"):
        app = AppTest.from_string(APP, default_timeout=20)
        app.secrets["app"] = {}
        app.run()
        yield app, storage
    view._records.clear(); view._json.clear(); view._live_job.clear(); view._quotation_folder.clear(); view._document_control_links.clear()


def test_new_screen_has_one_button_only_and_mask_updates_catalogue_without_scraping(ui):
    app, storage = ui
    assert not app.exception and not app.error
    assert [b.label for b in app.button] == ["Generar documentos"]
    assert not app.date_input and not app.radio and not app.tabs
    assert not any("Fechas" in x.label for x in app.expander)
    assert next(x for x in app.text_input if x.label == "Catálogo").value == "LB4330K"
    app.selectbox(key="anes_simple_mask").set_value("C").run()
    assert next(x for x in app.text_input if x.label == "Catálogo").value == "LB4330C"
    assert storage.enqueued == []


@pytest.mark.parametrize("quotation_only", [True, False])
def test_completion_shows_only_quotation_links_and_distinguishes_previous_full_packages(ui, quotation_only):
    app, storage = ui
    job = {"id": IDENT, "number": ACT, "url": URL, "state": "Documentos generados",
        "delivery_folder_id": "final-pdfs", "published_manifest": "ready-manifest",
        "final_url": "https://drive.google.com/drive/folders/final-pdfs",
        "word_url": "https://drive.google.com/file/d/word/view", "pdf_url": "https://drive.google.com/file/d/pdf/view", "zip_url": "https://drive.google.com/file/d/old-zip/view",
        "delivery_pdf_count": 1 if quotation_only else 12}
    if quotation_only:
        job["quotation_output_version"] = 3
    storage.tables["ANESTESIA_EXPEDIENTES"] = [job]
    storage.quotation_status = lambda job: True
    app.session_state["anes_simple_job"] = IDENT
    view._records.clear(); view._live_job.clear()
    app.run()
    assert not app.exception and not app.error
    links = "\n".join(m.value for m in app.markdown)
    assert "ZIP" not in links and "old-zip" not in links
    if quotation_only:
        assert any("Cotización membretada generada en PDF y Word" in s.value for s in app.success)
        assert "[Ver archivos en Drive]" in links and "[Cotización Word]" in links and "[Cotización PDF]" in links
    else:
        assert any("formato anterior" in i.value for i in app.info)
        assert not app.success


def test_quote_location_helper_imports_with_docs_module_from_previous_cloud_session(monkeypatch):
    import importlib
    from services import anestesia_docs
    monkeypatch.delattr(anestesia_docs, "validate_quotation")
    importlib.reload(quotes)
    assert quotes.location_fields({"purchase_unit": "Hospital", "info": {"provincia": "Panamá"}}) == ("Panamá", "Hospital")


def fill_ui(app, url=URL):
    app.text_input(key="anes_simple_url").set_value(url).run()
    app.number_input(key="anes_simple_price").set_value(19.48)
    next(w for w in app.checkbox if w.label.startswith("Revisé adjuntos")).check()
    next(w for w in app.checkbox if w.label.startswith("Verifiqué adjuntos")).check()
    app.button(key="anes_simple_generate").click().run()


def test_single_click_saves_inputs_and_queues_scrape_and_generate_together(ui):
    app, storage = ui
    fill_ui(app)
    assert not app.exception and not app.error
    assert len(storage.enqueued) == 1
    payload = storage.enqueued[0]
    assert payload["action"] == "generate_quotation" and payload["url"] == URL
    assert payload["config"]["price"] == "19.48"
    assert payload["config"]["delivery_use_portal"] and payload["config"]["location_reviewed"]
    assert next(b for b in app.button if b.label == "Generar documentos").disabled
    assert next(x for x in app.text_input if x.label == "Provincia").value == ""


def test_invalid_price_or_incomplete_delivery_does_not_enqueue(ui):
    app, storage = ui
    app.text_input(key="anes_simple_url").set_value(URL)
    app.button(key="anes_simple_generate").click().run()
    assert any("precio" in e.value for e in app.error) and not storage.enqueued
    app.number_input(key="anes_simple_price").set_value(20)
    app.button(key="anes_simple_generate").click().run()
    assert any("plazo" in e.value for e in app.error) and not storage.enqueued


def test_changed_link_does_not_inherit_adjunct_attestations_from_another_act(ui):
    app, storage = ui
    app.text_input(key="anes_simple_url").set_value(URL).run()
    next(w for w in app.checkbox if w.label.startswith("Revisé adjuntos")).check()
    next(w for w in app.checkbox if w.label.startswith("Verifiqué adjuntos")).check()
    app.run()
    assert all(w.value for w in app.checkbox)
    app.text_input(key="anes_simple_url").set_value(URL2).run()
    assert not any(w.value for w in app.checkbox)
    assert not storage.enqueued


def test_streamlit_hot_reload_replaces_old_storage_api_before_using_quotation_methods(monkeypatch):
    import importlib
    from services import anestesia_storage
    monkeypatch.setattr(anestesia_storage, "ANESTESIA_STORAGE_API_VERSION", 1)
    previous = anestesia_storage.AnestesiaStorage
    importlib.reload(view)
    assert view.AnestesiaStorage is anestesia_storage.AnestesiaStorage
    assert view.AnestesiaStorage is not previous
    assert anestesia_storage.ANESTESIA_STORAGE_API_VERSION == 3
    assert "ANESTESIA_COTIZACIONES" in anestesia_storage.TABLES
