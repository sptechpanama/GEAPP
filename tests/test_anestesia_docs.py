from copy import deepcopy
from datetime import datetime, timedelta
from io import BytesIO
import json
from pathlib import Path
import re
import zipfile

from docx import Document
import fitz
import pytest

from services.anestesia_docs import (AUDIT_CONTROLS, BASE_KINDS, COMPANY, PANAMA, canonical_hash,
    ANESTHESIA_WARRANTY, prepare_offer_config, document_status, file_hash, now_iso, public_registry_age, review_errors, totals, validate_package)
from services.anestesia_documents import pact_docx, quote_docx
from services.anestesia_source import delivery_destination, portal_delivery_term, tax_source_evidence, route, source_is_closed
from services import anestesia_worker as worker

ROOT = Path(__file__).resolve().parents[1]
TODAY = datetime.now(PANAMA).date()
IDENT = "a" * 32
ACT = "2026-1-10-01-08-CL-051191"
URL = "https://www.panamacompra.gob.pa/Inicio/#/solicitud-de-cotizacion/" + ACT + "/0nM6ICc0JCLwkjN2QDMxojIpJye"


def fixture():
    source = {"url": URL, "number": ACT, "title": "Kit de circuito de paciente para máquina de anestesia",
        "entity": "Caja de Seguro Social", "purchase_unit": "Ciudad de la Salud", "info": {"forma de pago": "Crédito"},
        "publication": str(TODAY - timedelta(days=10)), "closing": str(TODAY + timedelta(days=2)),
        "items": [{"cantidad": 900, "unidad": "Unidad"}], "explicit_fichas": ["43358"],
        "blocking_errors": [], "registry_max_months": 12, "registry_rule": "Anexo p.2, antigüedad máxima de un año",
        "fingerprint": "oficial-123"}
    config = {"catalog": "K", "price": "20", "tax_mode": "exento", "tax_rate": 7,
        "tax_evidence": "Anexo de presupuesto, p.1, ITBMS 0", "catalog_brand": "MFLab", "catalog_model": "LB4330K",
        "source_confirmed": True, "signature_authorized": True, "delivery_place": "Almacén de Ciudad de la Salud",
        "delivery": "300 unidades a 30 días, 300 a 45 días y 300 a 60 días calendario desde orden de compra",
        "warranty": "Esterilidad mínima de 24 meses desde cada entrega", "entity_representative": "Persona de Prueba",
        "entity_id": "8-000-000", "entity_role": "Director", "document_date": str(TODAY),
        "control_date": str(TODAY + timedelta(days=2)), "proposal_validity_days": 30}
    docs = [{"id": kind, "kind": kind, "company": COMPANY, "issued": str(TODAY - timedelta(days=5)),
        "expires": str(TODAY + timedelta(days=100)), "verified": True, "evidence": "Original completo, página 1",
        "catalogs": "K", "models": "LB4330K", "fichas": "43358", "act": ACT if kind in {"retorsion", "calidad"} else "",
        "notarized": True, "apostilled": True, "translation_verified": True,
        "file_id": kind, "sha256": "hash", "created_at": now_iso()} for kind in BASE_KINDS]
    return source, config, docs


@pytest.mark.parametrize("catalog,model", [("K", "LB4330K"), ("C", "LB4330C")])
def test_full_certificate_covers_both_variants_but_quote_selects_only_one(catalog, model):
    source, config, docs = fixture()
    config.update(catalog=catalog, catalog_model=model)
    for doc in docs:
        doc.update(catalogs="K,C", models="LB4330K, LB4330C", evidence="CT páginas 1 y 2; catálogo página 4")
    checks, chosen = validate_package(source, config, docs)
    assert all(c["estado"] == "Vigente documentalmente" for c in checks), checks
    assert chosen["criterio_tecnico"]["models"] == "LB4330K, LB4330C"


@pytest.mark.parametrize("catalog,model", [("C", "LB4330K"), ("K", "LB4330C")])
def test_full_certificate_does_not_allow_wrong_mask_selection(catalog, model):
    source, config, docs = fixture()
    config.update(catalog=catalog, catalog_model=model)
    for doc in docs: doc.update(catalogs="K,C", models="LB4330K,LB4330C")
    checks, _ = validate_package(source, config, docs)
    assert any("corresponde al catálogo" in c["motivo"] for c in checks)


@pytest.mark.parametrize("mode,subtotal,tax,total", [("exento", "18000.00", "0.00", "18000.00"),
    ("adicional", "18000.00", "1260.00", "19260.00"), ("incluido", "16822.43", "1177.57", "18000.00")])
def test_tax_never_double_counts_unit_price(mode, subtotal, tax, total):
    result = totals(900, "20", mode)
    assert (result["subtotal"], result["itbms"], result["total"]) == (subtotal, tax, total)


@pytest.mark.parametrize("price", [0, -5, "NaN", "inf", None, "1,000"])
def test_invalid_price_blocks_generation(price):
    with pytest.raises(ValueError): totals(900, price, "exento")


@pytest.mark.parametrize("text,months", [("Registro Público. Una vigencia no mayor de seis (6) meses", 6),
    ("Registro Público; vigencia no mayor de un (1) año", 12),
    ("Registro Público; vigencia no mayor de un (l) año", 12),
    ("Registro Público; antigüedad no superior a 3 meses", 3),
    ("Registro Público. El kit tiene vida útil de 24 meses", None)])
def test_registry_age_comes_from_this_tender(text, months):
    assert public_registry_age(text)[0] == months


def test_complete_evidence_passes_and_all_missing_documents_are_listed():
    source, config, docs = fixture()
    checks, _ = validate_package(source, config, docs)
    assert {r["estado"] for r in checks} == {"Vigente documentalmente"}
    checks, _ = validate_package(source, config, [])
    assert len(checks) == 11 and all(c["estado"] == "Falta" for c in checks)


def test_requirements_match_original_offers_and_reuse_legacy_imports():
    source, config, docs = fixture()
    expected = {'dgi', 'css', 'registro_publico', 'oferente', 'inscripcion_producto',
                'criterio_tecnico', 'catalogo', 'cedula', 'aviso_operacion', 'licencia_minsa', 'metodo_destruccion'}
    # Imported originals already exist in Drive under older classification names.
    next(d for d in docs if d['kind'] == 'licencia_minsa')['kind'] = 'otro:Licencia de operaciones MINSA'
    next(d for d in docs if d['kind'] == 'metodo_destruccion')['kind'] = 'disposicion'
    docs.extend([{'kind': kind, 'verified': False} for kind in ('retorsion', 'calidad')])
    before = deepcopy(docs)
    checks, selected = validate_package(source, config, docs)
    assert set(selected) == {c['kind'] for c in checks} == expected
    assert all(c['estado'] == 'Vigente documentalmente' for c in checks)
    assert selected['metodo_destruccion']['id'] == 'metodo_destruccion'
    assert docs == before


@pytest.mark.parametrize("kind,changes,reason", [
    ("css", {"expires": str(TODAY - timedelta(days=1)), "created_at": now_iso()}, "Venció"),
    ("dgi", {"expires": ""}, "vencimiento"),
    ("criterio_tecnico", {"catalogs": "C"}, "biblioteca"),
    ("criterio_tecnico", {"fichas": "102625"}, "43358"),
    ("oferente", {"company": "Otra empresa"}, "RIR"),
    ("licencia_minsa", {"expires": str(TODAY - timedelta(days=1))}, "Venció"),
    ("metodo_destruccion", {"expires": str(TODAY - timedelta(days=1))}, "Venció"),
    ("cedula", {"verified": "false"}, "verificar"),
    ("catalogo", {"file_id": ""}, "Drive"),
    ("registro_publico", {"issued": str(TODAY - timedelta(days=400))}, "meses"),
])
def test_certificates_fail_closed_with_specific_action(kind, changes, reason):
    source, config, docs = fixture()
    next(d for d in docs if d["kind"] == kind).update(changes)
    checks, _ = validate_package(source, config, docs)
    check = next(c for c in checks if c["kind"] == kind)
    assert check["estado"] != "Vigente documentalmente" and reason in check["motivo"]


def test_cannot_validate_only_today_when_certificate_expires_before_tender_close():
    source, config, docs = fixture()
    config["control_date"] = str(TODAY)
    docs[0]["expires"] = str(TODAY + timedelta(days=1))
    checks, _ = validate_package(source, config, docs)
    assert any("presentación" in c["motivo"] for c in checks)


def test_default_registry_policy_does_not_bypass_an_unreadable_annex():
    source, config, docs = fixture()
    source.update(registry_max_months=None, blocking_errors=["Anexo no interpretado"])
    checks, _ = validate_package(source, config, docs)
    assert any("Anexo no interpretado" in c["motivo"] for c in checks)
    assert next(c for c in checks if c['kind'] == 'registro_publico')['estado'] == 'Vigente documentalmente'


def test_old_saved_age_override_cannot_bypass_stricter_current_tender_at_generation():
    source, config, docs = fixture()
    source['registry_max_months'] = 3
    config.update(registry_max_months=12, registry_rule_evidence='Antiguo valor manual')
    next(d for d in docs if d['kind'] == 'registro_publico')['issued'] = str(TODAY - timedelta(days=150))
    checks, _ = validate_package(source, config, docs)
    check = next(c for c in checks if c['kind'] == 'registro_publico')
    assert check['estado'] == 'Bloqueado' and '3 meses' in check['motivo']


def test_explicit_additional_requirements_and_no_automatic_ct_for_c():
    source, config, docs = fixture()
    config.update(require_rs=True, require_power=True, extra_requirements=["Certificado adicional"])
    checks, _ = validate_package(source, config, docs)
    assert len(checks) == 14 and sum(c["estado"] == "Falta" for c in checks) == 3
    config["catalog"] = "C"
    assert any(c["kind"] == "criterio_tecnico" and c["estado"] == "Falta" for c in validate_package(source, config, docs)[0])


def test_latest_invalid_version_cannot_silently_fall_back_to_an_older_valid_one():
    source, config, docs = fixture()
    docs.append({**docs[0], "id": "updated", "verified": False, "created_at": "2099-01-01"})
    checks, selected = validate_package(source, config, docs)
    assert selected["dgi"]["id"] == "updated"
    assert next(c for c in checks if c["kind"] == "dgi")["estado"] == "Bloqueado"


def test_url_and_closing_reject_wrong_host_corrupt_tokens_and_past_or_cancelled_acts():
    assert route(URL) == (1046690, 2, ACT)
    for url in (URL.replace("panamacompra.gob.pa", "evil.example"), URL + "bad", "http://localhost/test"):
        with pytest.raises(ValueError): route(url)
    source, _, _ = fixture()
    assert not source_is_closed(source)[0]
    assert source_is_closed({**source, "closing": ""})[0]
    assert source_is_closed({**source, "closing": "2020-01-01"})[0]
    assert source_is_closed({**source, "info": {"estado": "Suspendido"}})[0]


def text_docx(data):
    doc = Document(BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_authored_documents_have_correct_amounts_delivery_identity_and_position():
    source, config, _ = fixture()
    quote = quote_docx(source, config, ROOT / "assets/cotizacion_base")
    text = text_docx(quote)
    assert "18,000.00" in text and "22,500.00" not in text
    for fragment in ("900", "LB4330K", "300 a 45", "Ciudad de la Salud", "24 meses", "S.EP."):
        assert fragment in text
    doc = Document(BytesIO(quote))
    assert len(doc.inline_shapes) >= 1
    assert "PH Bonanza" in "\n".join(p.text for s in doc.sections for p in s.header.paragraphs) or "PH Bonanza" in str(doc.sections[0].header._element.xml)
    pact = text_docx(pact_docx(source, config, ROOT))
    assert "en su calidad de Director" in pact and "Director" in pact.splitlines()
    assert "persona natural" not in pact and "sociedad de emprendimiento" in pact
    assert "155750585-2-2024-2024-574365876" in pact and "Bella Vista" in pact
    assert not re.search(r"\[[A-Za-z_]+\]", pact)


def test_quote_preserves_unit_price_precision_used_to_calculate_total():
    source, config, _ = fixture()
    config["price"] = "20.1234"
    quote = text_docx(quote_docx(source, config, ROOT / "assets/cotizacion_base"))
    assert "20.1234" in quote and "18,111.06" in quote


def approval(manifest):
    return {"request_id": manifest["request_id"], "manifest_hash": manifest["manifest_hash"],
        "decision": "aprobado", "pendientes": [], "revisor": "Revisor de prueba", "modelo": "simulacion",
        "reviewed_at": now_iso(), "archivos_revisados_sha256": [f["sha256"] for f in manifest["files"] if f.get("deliverable")],
        "controles": {key: {"resultado": "cumple", "evidencia": "Simulación automatizada; no constituye revisión de una oferta real"} for key in AUDIT_CONTROLS}}


@pytest.mark.parametrize("change", [{"manifest_hash": "wrong"}, {"decision": "observado"}, {"pendientes": ["Falta CT"]},
    {"controles": {}}, {"archivos_revisados_sha256": []}, {"archivos_revisados_sha256": [{}]}, {"reviewed_at": "2099-01-01"}])
def test_incomplete_or_other_version_review_cannot_approve(change):
    manifest = {"request_id": IDENT, "manifest_hash": "h", "created_at": now_iso(), "files": [{"deliverable": True, "sha256": "f"}]}
    assert not review_errors(approval(manifest), manifest)
    assert review_errors({**approval(manifest), **change}, manifest)


class FakeStorage:
    def __init__(self, source, docs):
        self.objects = {"source": json.dumps(source).encode()}
        self.writes = []
        self.saved = {"id": IDENT, "number": ACT, "state": "Datos capturados", "source_id": "source"}
        self.tables = {"ANESTESIA_DOCUMENTOS": deepcopy(docs), "ANESTESIA_REVISIONES": []}
        for item in self.tables["ANESTESIA_DOCUMENTOS"]:
            issuer = 'DIRECCION GENERAL DE INGRESOS' if item['kind'] == 'dgi' else 'CAJA DEL SEGURO SOCIAL\nNumero patronal: 123'
            with fitz.open() as pdf:
                pdf.new_page().insert_text((40, 40), f"{issuer}\nRIR MEDICAL ENGINEERING\nEmision: {item['issued']}\nValido hasta: {item['expires']}")
                data = pdf.tobytes()
            item["sha256"] = file_hash(data); self.objects[item["file_id"]] = data
    def job(self, ident): return dict(self.saved)
    def save_job(self, changes): self.saved.update(changes); return dict(self.saved)
    def root(self): return "root"
    def folder(self, name, parent): return name
    def rows(self, name): return deepcopy(self.tables[name])
    def json_file(self, ident): return json.loads(self.objects[ident])
    def get_bytes(self, ident): return self.objects[ident]
    def put(self, parent, name, data, mime):
        ident = str(len(self.objects)); self.objects[ident] = data
        self.writes.append({'parent': parent, 'name': name, 'mime': mime, 'file_id': ident})
        return {"file_id": ident, "name": name, "url": "https://drive.example/" + ident, "sha256": file_hash(data), "mime": mime}
    def convert_document(self, data, name, parent):
        pdf = fitz.open(); page = pdf.new_page(); page.insert_text((40, 40), ACT)
        return pdf.tobytes()


def generated(monkeypatch):
    source, config, docs = fixture()
    storage = FakeStorage(source, docs)
    monkeypatch.setattr(worker, "capture", lambda *args, **kw: source)
    class Publisher:
        def __init__(self, backend): self.backend = backend
        def publish(self, pdfs, **metadata):
            self.backend.published_pdfs = deepcopy(pdfs)
            return {'folder_id': 'stable-delivery-folder', 'count': len(pdfs)}
    monkeypatch.setattr(worker, 'DeliveryPublisher', Publisher)
    result = worker.run_request(storage, {"action": "generate", "request_id": IDENT, "config": config}, execution_id="first", root=ROOT)
    assert result["state"] == "Pendiente de revisión"
    manifest = storage.json_file(result["manifest_id"])
    review = {**approval(manifest), "id": "review"}
    storage.tables["ANESTESIA_REVISIONES"].append(review)
    return source, storage, manifest


def test_wrong_certificate_content_blocks_before_creating_quotation(monkeypatch):
    source, config, docs = fixture()
    storage = FakeStorage(source, docs)
    with fitz.open() as pdf:
        pdf.new_page().insert_text((40, 40), 'COTIZACION DE PRUEBA - NO ES UN PAZ Y SALVO')
        data = pdf.tobytes()
    storage.objects['css'] = data
    next(d for d in storage.tables['ANESTESIA_DOCUMENTOS'] if d['kind'] == 'css')['sha256'] = file_hash(data)
    before = set(storage.objects)
    monkeypatch.setattr(worker, 'capture', lambda *a, **kw: source)
    result = worker.run_request(storage, {'action':'generate','request_id':IDENT,'config':config}, execution_id='wrong-pdf',root=ROOT)
    assert result['state'] == 'Bloqueado'
    assert set(storage.objects) == before  # no quotation, pact or other deliverable created
    assert 'paz y salvo CSS' in result['checks'][0]['motivo']


def test_worker_generates_reviewable_bundle_and_only_publishes_exact_approved_files(monkeypatch):
    _, storage, manifest = generated(monkeypatch)
    assert len(manifest["files"]) == 13  # 12 PDF plus the editable quotation in the audit folder
    assert manifest['delivery_pdf_count'] == 12
    assert not any('Pacto' in f['name'] for f in manifest['files'])
    result = worker.run_request(storage, {"action": "finalize", "request_id": IDENT, "review_id": "review", "user_confirmed": True}, execution_id="second", root=ROOT)
    assert result["state"] == "Listo para entregar" and result["zip_url"]
    assert result['delivery_pdf_count'] == 12 and result['final_url'].endswith('stable-delivery-folder')
    assert len(storage.published_pdfs) == 12 and all(f['mime'] == 'application/pdf' for f in storage.published_pdfs)
    case_folder = result['archive_url'].split('/')[-1]
    assert ACT in case_folder and 'PDF para presentar' in case_folder
    case_files = [f for f in storage.writes if f['parent'] == case_folder]
    assert len(case_files) == 12 and all(f['mime'] == 'application/pdf' for f in case_files)
    assert sum(f['name'] == '01_Cotizacion.pdf' for f in case_files) == 1
    assert {f['name'] for f in case_files} == {f['name'] for f in storage.published_pdfs}
    archive_id = result['zip_url'].split('/')[-1]
    with zipfile.ZipFile(BytesIO(storage.objects[archive_id])) as zipped:
        assert len(zipped.namelist()) == 12 and all(name.endswith('.pdf') for name in zipped.namelist())
    size = len(storage.objects)
    again = worker.run_request(storage, {"action": "finalize", "request_id": IDENT}, execution_id="retry", root=ROOT)
    assert again["published_manifest"] == manifest["manifest_hash"] and len(storage.objects) == size


@pytest.mark.parametrize("tamper", ["pdf", "manifest", "source", "library", "expired", "closed", "approval"])
def test_publication_blocks_every_material_change_after_review(monkeypatch, tamper):
    source, storage, manifest = generated(monkeypatch)
    if tamper == "pdf": storage.objects[manifest["files"][0]["file_id"]] += b"modified"
    elif tamper == "manifest": storage.objects[storage.saved["manifest_id"]] = json.dumps({**manifest, "amounts": {"total": "1"}}).encode()
    elif tamper == "source": monkeypatch.setattr(worker, "capture", lambda *a, **k: {**source, "fingerprint": "changed"})
    elif tamper == "library": storage.tables["ANESTESIA_DOCUMENTOS"][0]["id"] = "changed"
    elif tamper == "expired": storage.tables["ANESTESIA_DOCUMENTOS"][0]["expires"] = "2020-01-01"
    elif tamper == "closed": monkeypatch.setattr(worker, "capture", lambda *a, **k: {**source, "closing": "2020-01-01"})
    elif tamper == "approval": storage.tables["ANESTESIA_REVISIONES"][0]["decision"] = "observado"
    try:
        result = worker.run_request(storage, {"action": "finalize", "request_id": IDENT, "review_id": "review", "user_confirmed": True}, execution_id="changed", root=ROOT)
        assert result["state"] == "Bloqueado"
    except ValueError:
        pass
    assert storage.saved["state"] != "Listo para entregar" and not storage.saved.get("zip_url")


def test_capture_error_preserves_previous_source_and_library(monkeypatch):
    source, _, docs = fixture()
    storage = FakeStorage(source, docs)
    before = deepcopy(storage.objects)
    def fail(*a, **k): raise TimeoutError("Fuente temporalmente caída")
    monkeypatch.setattr(worker, "capture", fail)
    with pytest.raises(TimeoutError):
        worker.run_request(storage, {"action": "capture", "request_id": IDENT, "url": URL}, execution_id="failed", root=ROOT)
    assert storage.objects == before and storage.saved["source_id"] == "source"


@pytest.mark.parametrize('catalog', ['K', 'C'])
def test_initial_capture_preserves_offer_entered_before_the_worker_runs(monkeypatch, catalog):
    source, _, docs = fixture()
    storage = FakeStorage(source, docs)
    config = {'catalog': catalog, 'price': '19.875', 'tax_mode': 'incluido', 'tax_rate': 7}
    storage.saved.update(state='En cola', config=deepcopy(config))
    monkeypatch.setattr(worker, 'capture', lambda *args: source)
    result = worker.run_request(storage, {'action': 'capture', 'request_id': IDENT, 'url': URL}, execution_id='initial-capture', root=ROOT)
    assert result['state'] == 'Datos capturados' and result['config'] == config
    assert storage.json_file(result['source_id']) == source
    assert not result.get('manifest_id')  # capture is not a generated or approved bid


@pytest.mark.parametrize('text,place', [
    ('LUGAR DE ENTREGA\nCIUDAD SALUD – ALMACEN GENERAL\n30 DÍAS CALENDARIO\nTIEMPO DE ENTREGA', 'CIUDAD SALUD – ALMACEN GENERAL'),
    ('DE\nLUGAR\nENTREGA\nHOSPITAL DR,G,N,C.R, ALM. MEDICOOUIRURGICO\nTIEMPO\nENTREGA\n30 DÍAS CALENDARIOS', 'HOSPITAL DR,G,N,C.R, ALM. MEDICOOUIRURGICO'),
    ('LUGAR DE ENTREGA:\r\nVIGENCIA:\r\n30 HABILES\r\nUNIDAD\r\nPANAMA J.J.VALLARINO.Z - ALMACEN GENERAL\r\nNO APLICA\r\nCTNI:', 'PANAMA J.J.VALLARINO.Z - ALMACEN GENERAL'),
    ('Lugar de entrega: Almacén general del hospital.\nTiempo de entrega: 30 días', 'Almacén general del hospital.'),
    ('LUGAR DE ENTREGA\nAlmacén general\nPoliclínica Joaquín José Vallarino\nPlanta baja\nTIEMPO DE ENTREGA\n30 DÍAS', 'Almacén general Policlínica Joaquín José Vallarino Planta baja'),
    ('Lugar de entrega: Almacén general. Tiempo de entrega: 30 días', 'Almacén general.'),
])
def test_destination_comes_from_official_annex_including_reordered_ocr_tables(text, place):
    source = {'info': {'provincia de entrega': 'Panamá', 'direccion de la unidad de compra': 'Juan Díaz'},
              'attachments': [{'name': 'Requerimientos.pdf', 'text': text}]}
    found, evidence = delivery_destination(source)
    assert found == place and 'Requerimientos.pdf' in evidence


@pytest.mark.parametrize('text', ['', 'LUGAR DE ENTREGA\nVIGENCIA\n24 MESES\nUNIDAD\nCTNI: 43358',
                                  'Requisitos: cotización dirigida al hospital de prueba.'])
def test_destination_never_invents_warehouse_from_buyer_or_province(text):
    source = {'purchase_unit': 'Hospital de prueba', 'info': {'provincia de entrega': 'Panamá'},
              'attachments': [{'text': text}]}
    assert delivery_destination(source)[0] == ''


def test_explicit_portal_destination_wins_over_unrelated_annex():
    source = {'info': {'Lugar de entrega': 'Almacén de Vallarino'},
              'attachments': [{'text': 'LUGAR DE ENTREGA: Almacén antiguo'}]}
    assert delivery_destination(source) == ('Almacén de Vallarino', 'PanamáCompra: Lugar de entrega')


def test_offer_defaults_remove_generic_pact_fields_and_use_standing_signature_permission():
    source, cfg, docs = fixture()
    source['info']['lugar de entrega'] = 'Almacén médico del acto'
    cfg.update(delivery_place='Destino anterior', require_rs=True, require_power=True, signature_authorized=False)
    before = deepcopy(cfg)
    result = prepare_offer_config(source, cfg)
    assert cfg == before
    assert result['delivery_place'] == 'Almacén médico del acto'
    assert result['warranty'] == ANESTHESIA_WARRANTY
    assert not result['require_rs'] and not result['require_power']
    assert result['signature_authorized'] is True
    assert not any(field in result for field in ['entity_representative', 'entity_id', 'entity_role'])
    checks, chosen = validate_package(source, result, docs)
    assert all(c['estado'] == 'Vigente documentalmente' for c in checks), checks
    assert not {'poder', 'registro_sanitario'} & set(chosen)


def test_stricter_explicit_sterility_requirement_cannot_be_replaced_silently_by_default():
    source, cfg, docs = fixture()
    source['items'][0]['descripcion'] = 'Vencimiento de la esterilidad no menor de 36 meses a partir de la entrega.'
    checks, _ = validate_package(source, prepare_offer_config(source, cfg), docs)
    assert any('36 meses' in c['motivo'] for c in checks)


def test_worker_generates_own_signed_quote_without_requesting_entity_or_proxy_data(monkeypatch):
    source, cfg, docs = fixture()
    for field in ['entity_representative', 'entity_id', 'entity_role', 'warranty', 'signature_authorized']:
        cfg.pop(field, None)
    source['info']['lugar de entrega'] = 'Almacén general de la unidad solicitante'
    cfg.pop('delivery_place')
    storage = FakeStorage(source, docs)
    monkeypatch.setattr(worker, 'capture', lambda *args, **kwargs: source)
    result = worker.run_request(storage, {'action': 'generate', 'request_id': IDENT, 'config': cfg}, execution_id='simplified-offer', root=ROOT)
    assert result['state'] == 'Pendiente de revisión'
    manifest = storage.json_file(result['manifest_id'])
    quote = next(f for f in manifest['files'] if f['name'] == '01_Cotizacion.docx')
    content = text_docx(storage.get_bytes(quote['file_id']))
    assert '24 meses' in content and 'Almacén general de la unidad solicitante' in content
    assert len(Document(BytesIO(storage.get_bytes(quote['file_id']))).inline_shapes) >= 1
    assert result['config']['signature_authorized'] is True


@pytest.mark.parametrize('key', ['termino de entrega', 'Término de entrega', 'Plazo de entrega', 'Tiempo de entrega:'])
def test_portal_delivery_reads_only_delivery_period_and_preserves_day_type(key):
    assert portal_delivery_term({'info': {key: '30 Días hábiles', 'fecha de entrega de propuestas': '2026-10-01'}}) == '30 Días hábiles'
    assert portal_delivery_term({'info': {'fecha de entrega de propuestas': '2026-10-01', 'forma de entrega': 'Total'}}) == ''


@pytest.mark.parametrize('use_portal', [True, False])
def test_explicit_delivery_choice_controls_offer_without_overwriting_manual_schedule(use_portal):
    source, cfg, docs = fixture()
    source['info']['termino de entrega'] = '60 Días calendario'
    manual = '300 unidades a 30 días; 300 a 45 días; 300 a 60 días calendario desde la orden de compra.'
    cfg.update(delivery_use_portal=use_portal, delivery_manual=manual,
               delivery_portal_value='60 Días calendario', delivery_portal_source_hash=source['fingerprint'])
    result = prepare_offer_config(source, cfg)
    assert result['delivery'] == ('60 Días calendario' if use_portal else manual)
    assert result['delivery_manual'] == manual
    checks, _ = validate_package(source, result, docs)
    assert all(c['estado'] == 'Vigente documentalmente' for c in checks), checks


def test_portal_is_never_used_without_check_when_manual_schedule_is_empty():
    source, cfg, docs = fixture()
    source['info']['termino de entrega'] = '30 Días hábiles'
    cfg.update(delivery_use_portal=False, delivery_manual='')
    result = prepare_offer_config(source, cfg)
    assert result['delivery'] == ''
    checks, _ = validate_package(source, result, docs)
    assert any('calendario completo' in c['motivo'] for c in checks)


@pytest.mark.parametrize('change', ['missing', 'changed_period', 'changed_annex', 'unconfirmed'])
def test_portal_delivery_requires_current_review_and_cannot_hide_missing_period(change):
    source, cfg, docs = fixture()
    source['info']['termino de entrega'] = '30 Días hábiles'
    cfg.update(delivery_use_portal=True, delivery_portal_value='30 Días hábiles', delivery_portal_source_hash=source['fingerprint'])
    if change == 'missing': source['info'].pop('termino de entrega')
    elif change == 'changed_period': source['info']['termino de entrega'] = '45 Días calendario'
    elif change == 'changed_annex': source['fingerprint'] = 'modified-annex'
    elif change == 'unconfirmed': cfg.pop('delivery_portal_value')
    checks, _ = validate_package(source, prepare_offer_config(source, cfg), docs)
    assert any(c.get('kind') == 'expediente' and c['estado'] == 'Bloqueado' for c in checks)


@pytest.mark.parametrize('catalog,model', [('K', 'LB4330K'), ('C', 'LB4330C')])
def test_model_follows_mask_even_when_previous_configuration_had_the_other_model(catalog, model):
    source, cfg, _ = fixture()
    cfg.update(catalog=catalog, catalog_brand='Old brand', catalog_model='Old model')
    result = prepare_offer_config(source, cfg)
    assert result['catalog_model'] == model and result['catalog_brand'] == 'MFLAB'


@pytest.mark.parametrize('mode', ['exento', 'adicional', 'incluido'])
def test_tax_evidence_not_mandatory_and_does_not_override_selected_tax_mode(mode):
    source, cfg, docs = fixture()
    cfg.pop('tax_evidence')
    cfg['tax_mode'] = mode
    result = prepare_offer_config(source, cfg)
    checks, _ = validate_package(source, result, docs)
    assert all(c['estado'] == 'Vigente documentalmente' for c in checks), checks
    assert result['tax_evidence'] == ''
    assert result['tax_mode'] == mode
    source['items'][0]['itbms'] = 0
    result = prepare_offer_config(source, cfg)
    assert result['tax_mode'] == mode and result['tax_rate'] == cfg['tax_rate']
    assert result['tax_evidence_auto'] == [{'fuente': 'PanamáCompra, renglón 1: ITBMS', 'valor': '0'}]


def test_available_tax_observations_and_legacy_notes_are_preserved_without_inventing_evidence():
    source, cfg, _ = fixture()
    source['attachments'] = [{'name': 'Presupuesto.pdf', 'text': 'ITBMS: 0.00%\nOtra información'}]
    result = prepare_offer_config(source, cfg)
    assert result['tax_evidence'] == cfg['tax_evidence']
    assert tax_source_evidence(source) == [{'fuente': 'Presupuesto.pdf', 'valor': 'ITBMS: 0.00%'}]


def test_worker_quote_uses_confirmed_portal_term_and_no_tax_note_is_required(monkeypatch):
    source, cfg, docs = fixture()
    source['info']['termino de entrega'] = '30 Días hábiles'
    cfg.update(delivery_use_portal=True, delivery_manual='300 unidades a 45 días',
               delivery_portal_value='30 Días hábiles', delivery_portal_source_hash=source['fingerprint'])
    cfg.pop('tax_evidence')
    storage = FakeStorage(source, docs)
    monkeypatch.setattr(worker, 'capture', lambda *args, **kwargs: source)
    result = worker.run_request(storage, {'action': 'generate', 'request_id': IDENT, 'config': cfg}, execution_id='portal-delivery', root=ROOT)
    assert result['state'] == 'Pendiente de revisión'
    manifest = storage.json_file(result['manifest_id'])
    quote = next(f for f in manifest['files'] if f['name'] == '01_Cotizacion.docx')
    content = text_docx(storage.get_bytes(quote['file_id']))
    assert 'Entregas: 30 Días hábiles' in content and '300 unidades a 45 días' not in content
